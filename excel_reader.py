"""唯讀讀取 Excel:一律先複製到暫存,讀完立即刪除,絕不碰原始檔。"""

from __future__ import annotations

import io
import logging
import re
import shutil
import tempfile
import math
import posixpath
import stat as stat_module
import threading
import time
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string

from cell_format import display_value, normalize, sort_key, split_keywords
from table_layout import Merge, build_table, drop_hidden

logger = logging.getLogger(__name__)

# 加密的 .xlsx 與舊版 .xls 都是 OLE 複合文件(不是 zip)
OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
# 加密檔的 OLE 目錄中會有這個串流名稱(UTF-16LE)
ENCRYPTED_MARK = "EncryptedPackage".encode("utf-16-le")

MSG_ENCRYPTED = "此檔案有密碼保護,無法讀取"
MSG_LEGACY_XLS = "此檔案是舊版 .xls 格式,請另存為 .xlsx"
MSG_CORRUPTED = "此檔案已損毀,無法讀取"
MSG_UNAVAILABLE = "此檔案目前無法存取,請稍後再試"


ROW_SEPARATOR = "\x1f"  # 組合搜尋字串時分隔儲存格,避免跨儲存格誤判為符合


class ExcelReadError(Exception):
    """讀取失敗;訊息是給使用者看的簡短原因。"""


class SheetNotFound(Exception):
    """指定的工作表不存在。"""


class SortColumnNotFound(Exception):
    """指定的排序欄位不存在。"""


@dataclass
class Sheet:
    """一個工作表:欄位名稱、顯示用的各列文字,以及每列正規化後的搜尋字串。"""

    columns: list[str]
    rows: list[list[str]]
    haystack: pd.Series
    # 各欄遞增排序後的列索引(空白排最後),第一次排序時計算後保留
    _sort_cache: dict[int, tuple[np.ndarray, np.ndarray]] = field(default_factory=dict, repr=False, compare=False)

    @classmethod
    def build(cls, columns: list[str], rows: list[list[str]]) -> Sheet:
        haystack = pd.Series(
            [normalize(ROW_SEPARATOR.join(row)) for row in rows], dtype=object
        )
        return cls(columns, rows, haystack)

    def match_mask(self, keywords: list[str]) -> np.ndarray:
        """每列是否包含所有已正規化的關鍵字(AND,可分別在不同儲存格);沒有關鍵字時全部符合。"""
        mask = np.ones(len(self.rows), dtype=bool)
        for keyword in keywords:
            mask &= self.haystack.str.contains(keyword, regex=False).to_numpy(dtype=bool)
        return mask

    def sorted_order(self, column: int, descending: bool = False) -> np.ndarray:
        """依某欄排序後的列索引。數字依數值、文字依字串;不論遞增遞減,空白都排在最後。"""
        if column not in self._sort_cache:
            filled = [i for i, row in enumerate(self.rows) if row[column]]
            blank = [i for i, row in enumerate(self.rows) if not row[column]]
            filled.sort(key=lambda i: sort_key(self.rows[i][column]))
            self._sort_cache[column] = (np.array(filled, dtype=np.int64), np.array(blank, dtype=np.int64))
        filled, blank = self._sort_cache[column]
        return np.concatenate([filled[::-1] if descending else filled, blank])


EXCEL_SUFFIXES = (".xlsx", ".xlsm")  # 副檔名不分大小寫;子資料夾不讀取


def is_excel_file(path: Path) -> bool:
    """是否為要讀取的 Excel 檔(.xlsx、.xlsm,不分大小寫),排除 Excel 開啟中的鎖定檔 ~$*。"""
    return path.suffix.lower() in EXCEL_SUFFIXES and not path.name.startswith("~$")


@dataclass(frozen=True)
class FileStat:
    path: Path
    mtime: float
    size: int


def scan_folder(data_dir: Path) -> dict[str, FileStat]:
    """列出資料夾中的 Excel(分類名稱 → 路徑、修改時間、大小),依名稱排序。

    資料夾不存在或無法存取(例如網路共用資料夾斷線)時拋出 OSError。
    """
    if not data_dir.is_dir():
        raise FileNotFoundError(f"找不到資料夾:{data_dir}")
    found: list[FileStat] = []
    for path in data_dir.iterdir():
        if not is_excel_file(path):
            continue
        try:
            info = path.stat()
        except OSError:  # 列出後被刪除或改名,下次再處理
            continue
        if stat_module.S_ISREG(info.st_mode):
            found.append(FileStat(path, info.st_mtime, info.st_size))
    # 分類名稱為檔名去掉副檔名;同名不同副檔名(報表.xlsx 與 報表.xlsm)時改用完整檔名區分
    stems = [f.path.stem for f in found]
    files = {(f.path.stem if stems.count(f.path.stem) == 1 else f.path.name): f for f in found}
    return dict(sorted(files.items()))


def list_categories(data_dir: Path) -> list[str]:
    """回傳資料夾中所有 Excel 的分類名稱;資料夾無法存取時回傳空清單。"""
    try:
        return list(scan_folder(data_dir))
    except OSError:
        return []


def _describe_failure(tmp_path: Path) -> str:
    """依暫存複本的內容判斷讀取失敗的原因。"""
    data = tmp_path.read_bytes()
    if data.startswith(OLE_MAGIC):
        return MSG_ENCRYPTED if ENCRYPTED_MARK in data else MSG_LEGACY_XLS
    return MSG_CORRUPTED


# 工作表 XML 中的 <mergeCell ref="A3:A5"/>(可能帶命名空間前綴)
_MERGE_REF = re.compile(rb'<(?:\w+:)?mergeCell\b[^>]*?\bref="([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?"')
TYPE_SAMPLE_ROWS = 60  # 只記錄前 60 列的儲存格型別,供偵測表頭使用


# openpyxl 唯讀工作表記錄 XML 位置的內部屬性(非公開 API,升級 openpyxl 後可能改名)
SHEET_PATH_ATTR = "_worksheet_path"
MSG_MERGE_UNSUPPORTED = (
    "目前的 openpyxl 版本取不到合併儲存格範圍(內部屬性 %s 不存在),"
    "合併儲存格將不會填入值;請改回 requirements.txt 鎖定的版本"
)
_merge_warning_logged = False


_ROW_TAG = re.compile(rb"<(?:\w+:)?row\b([^>]*)>")
_COL_TAG = re.compile(rb"<(?:\w+:)?col\b([^>]*)>")
_AUTO_FILTER = re.compile(rb'<(?:\w+:)?autoFilter\b[^>]*?\bref="[A-Z]+(\d+)(?::[A-Z]+(\d+))?"')
_RELATIONSHIP = re.compile(rb"<(?:\w+:)?Relationship\b([^>]*)>")


def _attr(attrs: bytes, name: str) -> str | None:
    match = re.search(rb"\b" + name.encode() + rb'="([^"]*)"', attrs)
    return match.group(1).decode() if match else None


def _sheet_xml(archive: zipfile.ZipFile, ws) -> tuple[str, bytes] | None:
    """取得工作表在檔案中的路徑與 XML。openpyxl 唯讀模式不提供合併範圍與隱藏列欄,只能自行解析。

    取不到時回傳 None(不當掉),並記錄一次警告。
    """
    global _merge_warning_logged
    sheet_path = getattr(ws, SHEET_PATH_ATTR, None)
    if not sheet_path:
        if not _merge_warning_logged:
            logger.warning(MSG_MERGE_UNSUPPORTED, SHEET_PATH_ATTR)
            _merge_warning_logged = True
        return None
    sheet_path = sheet_path.lstrip("/")
    try:
        return sheet_path, archive.read(sheet_path)
    except KeyError:
        return None


def _parse_merges(xml: bytes) -> list[Merge]:
    if b"mergeCell" not in xml:
        return []
    merges = []
    for c1, r1, c2, r2 in _MERGE_REF.findall(xml):
        c2, r2 = c2 or c1, r2 or r1
        merges.append(Merge(
            int(r1) - 1, column_index_from_string(c1.decode()) - 1,
            int(r2) - 1, column_index_from_string(c2.decode()) - 1,
        ))
    return merges


def _merge_ranges(archive: zipfile.ZipFile, ws) -> list[Merge]:
    """讀取工作表的合併範圍;取不到時回傳空清單。"""
    found = _sheet_xml(archive, ws)
    return _parse_merges(found[1]) if found else []


def _filter_row_ranges(archive: zipfile.ZipFile, sheet_path: str, xml: bytes) -> list[tuple[int, int]]:
    """篩選範圍涵蓋的列(1 起算,含兩端):工作表的 autoFilter,以及「格式化為表格」各表格的 autoFilter。"""
    ranges = [(int(r1), int(r2 or r1)) for r1, r2 in _AUTO_FILTER.findall(xml)]
    if b"tablePart" not in xml:
        return ranges
    folder = posixpath.dirname(sheet_path)
    rels_path = posixpath.join(folder, "_rels", posixpath.basename(sheet_path) + ".rels")
    try:
        rels = archive.read(rels_path)
    except KeyError:
        return ranges
    for attrs in _RELATIONSHIP.findall(rels):
        if not (_attr(attrs, "Type") or "").endswith("/table"):
            continue
        target = _attr(attrs, "Target") or ""
        table_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
        try:
            table_xml = archive.read(table_path)
        except KeyError:
            continue
        ranges += [(int(r1), int(r2 or r1)) for r1, r2 in _AUTO_FILTER.findall(table_xml)]
    return ranges


def _parse_hidden(archive: zipfile.ZipFile, sheet_path: str, xml: bytes) -> tuple[set[int], list[tuple[int, int]]]:
    """回傳 (要略過的列, 要略過的欄範圍),皆為 0 起算。

    手動隱藏的列與欄略過;位於篩選範圍內的隱藏列是篩選造成的,照常讀取。
    """
    if b'hidden="1"' not in xml and b'hidden="true"' not in xml:
        return set(), []
    hidden_cols = []
    for attrs in _COL_TAG.findall(xml):
        if _attr(attrs, "hidden") in ("1", "true"):
            first, last = int(_attr(attrs, "min") or 0), int(_attr(attrs, "max") or 0)
            if first:
                hidden_cols.append((first - 1, (last or first) - 1))
    hidden_rows = []
    row_number = 0
    for attrs in _ROW_TAG.findall(xml):
        r = _attr(attrs, "r")
        row_number = int(r) if r else row_number + 1  # r 可省略,省略時為上一列的下一列
        if _attr(attrs, "hidden") in ("1", "true"):
            hidden_rows.append(row_number)
    if not hidden_rows:
        return set(), hidden_cols
    filters = _filter_row_ranges(archive, sheet_path, xml)
    manual = {r - 1 for r in hidden_rows if not any(first <= r <= last for first, last in filters)}
    return manual, hidden_cols


def merge_support_ok() -> bool:
    """在記憶體中建立含合併儲存格的活頁簿,確認目前的 openpyxl 取得到合併範圍。"""
    wb = Workbook()
    wb.active.merge_cells("A1:B2")
    buf = io.BytesIO()
    wb.save(buf)
    ro = load_workbook(buf, read_only=True)
    try:
        if not getattr(ro.worksheets[0], SHEET_PATH_ATTR, None):
            return False
        with zipfile.ZipFile(buf) as archive:
            return _merge_ranges(archive, ro.worksheets[0]) == [Merge(0, 0, 1, 1)]
    finally:
        ro.close()


def _read_sheet(
    ws, merges: list[Merge], hidden_rows: set[int] | None = None, hidden_cols: list[tuple[int, int]] | None = None
) -> Sheet:
    """讀出顯示文字的網格,移除隱藏的列與欄後,交給 build_table 偵測表頭、填入合併格、整理欄位。"""
    ws.reset_dimensions()  # 有些程式產生的檔案記錄的範圍不正確,改為實際讀到哪算到哪
    grid: list[list[str]] = []
    is_text: list[list[bool]] = []
    for row in ws.iter_rows():
        grid.append([display_value(cell.value, cell.number_format) for cell in row])
        if len(is_text) < TYPE_SAMPLE_ROWS:
            is_text.append([isinstance(cell.value, str) and bool(cell.value.strip()) for cell in row])
    width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in grid), default=0)
    width = max([width] + [m.max_col + 1 for m in merges if m.min_row < len(grid)])
    grid = [(r + [""] * width)[:width] for r in grid]
    is_text = [(r + [False] * width)[:width] for r in is_text]
    merges = [m for m in merges if m.min_row < len(grid)]
    if hidden_rows or hidden_cols:
        cols = {c for first, last in hidden_cols or [] for c in range(first, min(last, width - 1) + 1)}
        rows = {r for r in hidden_rows or set() if r < len(grid)}
        grid, is_text, merges = drop_hidden(grid, is_text, merges, rows, cols)
    columns, rows = build_table(grid, is_text, merges)
    return Sheet.build(columns, rows)


def _read_workbook(path: Path, include_hidden: bool, include_hidden_rows_columns: bool) -> dict[str, Sheet]:
    wb = load_workbook(path, read_only=True, data_only=True)  # data_only:公式讀取 Excel 存下的計算結果
    try:
        with zipfile.ZipFile(path) as archive:
            sheets = {}
            for ws in wb.worksheets:
                if not include_hidden and ws.sheet_state != "visible":  # 隱藏、極度隱藏的工作表預設不讀
                    continue
                found = _sheet_xml(archive, ws)  # 每個工作表的 XML 只讀一次
                merges, hidden_rows, hidden_cols = [], set(), []
                if found:
                    merges = _parse_merges(found[1])
                    if not include_hidden_rows_columns:
                        hidden_rows, hidden_cols = _parse_hidden(archive, *found)
                sheets[ws.title] = _read_sheet(ws, merges, hidden_rows, hidden_cols)
            return sheets
    finally:
        wb.close()  # 先釋放檔案,Windows 才能刪除暫存檔


def read_excel_safely(
    path: Path, include_hidden: bool = False, include_hidden_rows_columns: bool = False
) -> dict[str, Sheet]:
    """複製到暫存資料夾後讀取工作表,讀完立即刪除暫存檔。

    include_hidden 為 False 時不讀取隱藏與極度隱藏的工作表;include_hidden_rows_columns 為 False 時
    不讀取手動隱藏的列與欄(篩選造成的隱藏列照常讀取)。

    讀取失敗時拋出 ExcelReadError,訊息為簡短的中文原因。
    """
    with tempfile.TemporaryDirectory(prefix="excel-lookup-") as tmp:
        tmp_path = Path(tmp) / (path.stem + path.suffix.lower())  # 副檔名轉小寫,openpyxl 依副檔名判斷格式
        try:
            shutil.copyfile(path, tmp_path)
        except OSError as exc:
            raise ExcelReadError(MSG_UNAVAILABLE) from exc
        try:
            return _read_workbook(tmp_path, include_hidden, include_hidden_rows_columns)
        except Exception as exc:
            reason = _describe_failure(tmp_path)
            logger.debug("無法讀取 %s:%s(%r)", path.name, reason, exc)  # 警告由 Catalog 在狀態改變時記錄
            raise ExcelReadError(reason) from exc
        finally:
            tmp_path.unlink(missing_ok=True)


@dataclass
class _Entry:
    key: tuple[float, int]  # 讀取時檔案的修改時間與大小,用來偵測變動
    sheets: dict[str, Sheet]
    error: str | None = None


class Catalog:
    """資料夾中所有 Excel 的讀取結果。

    refresh() 由背景執行緒定期呼叫,讀取新增或變動的檔案;搜尋只使用已載入的資料,
    不會在搜尋時讀檔。檔案變動後重新載入完成前,搜尋仍使用舊資料。
    請求(categories、search 等)完全不存取資料夾:網路資料夾斷線時存取可能卡住數秒。
    資料夾暫時無法連線時保留已載入的資料,恢復後自動更新。
    """

    def __init__(self, data_dir: Path, include_hidden: bool = False, include_hidden_rows_columns: bool = False) -> None:
        self.data_dir = data_dir
        self.include_hidden = include_hidden
        self.include_hidden_rows_columns = include_hidden_rows_columns
        self._entries: dict[str, _Entry] = {}
        self._files: dict[str, FileStat] = {}  # 最後一次成功列出的檔案
        self._refresh_lock = threading.Lock()
        self.folder_ok = False
        self.last_success: float | None = None  # 最後一次成功讀取資料夾的時間
        self._offline_logged = False
        try:
            self._files = scan_folder(data_dir)  # 先列出檔名,載入完成前就能顯示「載入中」
            self.folder_ok = True
        except OSError:
            pass

    def has(self, name: str) -> bool:
        return name in self._files

    def folder_status(self) -> dict:
        """資料夾連線狀態,以及最後一次成功讀取的時間。"""
        last = datetime.fromtimestamp(self.last_success).strftime("%Y-%m-%d %H:%M") if self.last_success else None
        return {"ok": self.folder_ok, "last_success": last}

    def refresh(self) -> None:
        """載入新增或變動的檔案(小檔優先),移除已刪除的檔案。

        讀取失敗的檔案每次都重讀(可能只是暫時被鎖定或網路不穩),同樣的錯誤只記錄一次。
        資料夾無法連線時不更動已載入的資料。
        """
        with self._refresh_lock:
            try:
                files = scan_folder(self.data_dir)
            except OSError as exc:
                self.folder_ok = False
                if not self._offline_logged:
                    logger.warning("Excel 資料夾無法連線,保留目前已載入的資料:%s", exc)
                    self._offline_logged = True
                return
            if self._offline_logged:
                logger.info("Excel 資料夾已恢復連線")
                self._offline_logged = False
            self.folder_ok = True
            self._files = files
            for gone in set(self._entries) - set(files):
                del self._entries[gone]
            pending = []
            for name, info in files.items():
                key = (info.mtime, info.size)
                entry = self._entries.get(name)
                if entry is None or entry.key != key or entry.error:
                    pending.append((info.size, name, key))
            for _, name, key in sorted(pending):
                path = files[name].path
                previous = self._entries.get(name)
                try:
                    sheets = read_excel_safely(path, self.include_hidden, self.include_hidden_rows_columns)
                    self._entries[name] = _Entry(key, sheets)
                except ExcelReadError as exc:
                    self._entries[name] = _Entry(key, {}, str(exc))
                    if previous is None or previous.error != str(exc):
                        logger.warning("無法讀取 %s:%s", path.name, exc)
                else:
                    if previous is not None and previous.error:
                        logger.info("已可讀取 %s", path.name)
            self.last_success = time.time()

    def _summary(self, name: str) -> dict:
        """單一檔案的狀態摘要:status 為 ok、error(附 error 原因)或 loading(尚未載入)。"""
        entry = self._entries.get(name)
        info = self._files.get(name)
        mtime = entry.key[0] if entry else (info.mtime if info else None)
        summary = {
            "name": name,
            "status": "loading" if entry is None else ("error" if entry.error else "ok"),
            "error": entry.error if entry else None,
            "updated": datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M") if mtime else None,
            "rows": None,
            "sheets": None,
        }
        if entry is not None and not entry.error:
            summary["rows"] = sum(len(sheet.rows) for sheet in entry.sheets.values())
            summary["sheets"] = len(entry.sheets)
        return summary

    def categories(self) -> list[dict]:
        """檔案清單:名稱、狀態、筆數、工作表數、最後更新時間(檔案修改時間)。"""
        return [self._summary(name) for name in self._files]

    def file_info(self, name: str) -> dict:
        """單一檔案的摘要與工作表清單;檔案不存在時拋出 KeyError。"""
        if name not in self._files:
            raise KeyError(name)
        info = self._summary(name)
        entry = self._entries.get(name)
        info["sheet_list"] = (
            [{"name": n, "rows": len(sh.rows), "columns": len(sh.columns)} for n, sh in entry.sheets.items()]
            if entry is not None and not entry.error
            else []
        )
        return info

    def rows(
        self,
        name: str,
        sheet: str | None = None,
        query: str = "",
        sort: int | None = None,
        descending: bool = False,
        page: int = 1,
        page_size: int = 100,
    ) -> dict:
        """某工作表的分頁資料,可篩選與排序。

        檔案不存在拋出 KeyError;工作表不存在拋出 SheetNotFound;排序欄位不存在拋出 SortColumnNotFound。
        頁碼超過範圍時改為最後一頁。檔案尚未載入或讀取失敗時 status 為 loading / error,rows 為空。
        """
        info = self.file_info(name)
        result = {
            "status": info["status"], "error": info["error"], "sheet": None, "columns": [], "rows": [],
            "total": 0, "matched": 0, "page": 1, "page_size": page_size, "pages": 1, "offset": 0,
        }
        if info["status"] != "ok":
            return result
        sheets = self._entries[name].sheets
        if not sheets:
            return result
        sheet_name = sheet if sheet is not None else next(iter(sheets))
        if sheet_name not in sheets:
            raise SheetNotFound(sheet_name)
        data = sheets[sheet_name]
        if sort is not None and not 0 <= sort < len(data.columns):
            raise SortColumnNotFound(sort)
        mask = data.match_mask(split_keywords(query))
        if sort is None:
            order = np.flatnonzero(mask)
        else:
            order = data.sorted_order(sort, descending)
            order = order[mask[order]]
        matched = len(order)
        pages = max(1, math.ceil(matched / page_size))
        page = min(max(page, 1), pages)
        offset = (page - 1) * page_size
        result.update(
            sheet=sheet_name,
            columns=data.columns,
            rows=[data.rows[i] for i in order[offset : offset + page_size]],
            total=len(data.rows),
            matched=matched,
            page=page,
            pages=pages,
            offset=offset,
        )
        return result

    def search(self, query: str, category: str | None = None, max_results: int = 500) -> dict:
        """在指定分類(或全部分類)中搜尋包含關鍵字的列。

        以空白拆成多個關鍵字,每個都要出現在同一列;不分大小寫、/ 與 - 視為相同、忽略數字間的千分位逗號。
        讀不了的檔案列在 errors,尚未載入完成的檔案列在 loading,都不影響其他檔案。
        """
        keywords = split_keywords(query)
        results: list[dict] = []
        errors: list[dict] = []
        loading: list[str] = []
        if not keywords:
            return {"results": results, "errors": errors, "loading": loading}
        remaining = max_results
        for cat in [category] if category else list(self._files):
            entry = self._entries.get(cat)
            if entry is None:
                loading.append(cat)
                continue
            if entry.error:
                errors.append({"category": cat, "error": entry.error})
                continue
            for sheet_name, sheet in entry.sheets.items():
                if remaining <= 0:
                    break
                if not sheet.rows:
                    continue
                indexes = np.flatnonzero(sheet.match_mask(keywords))[:remaining]
                if len(indexes) == 0:
                    continue
                remaining -= len(indexes)
                results.append(
                    {
                        "category": cat,
                        "sheet": sheet_name,
                        "columns": sheet.columns,
                        "rows": [sheet.rows[i] for i in indexes],
                    }
                )
        return {"results": results, "errors": errors, "loading": loading}
