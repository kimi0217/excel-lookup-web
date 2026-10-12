from datetime import date, datetime, time
from pathlib import Path

import pytest
from openpyxl import Workbook

import excel_reader
from cell_format import display_value, format_number, normalize
from conftest import search


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (datetime(2026, 1, 15), "2026-01-15"),
        (datetime(2026, 12, 31, 17, 30), "2026-12-31 17:30"),
        (datetime(2026, 12, 31, 17, 30, 5), "2026-12-31 17:30:05"),
        (date(2026, 3, 1), "2026-03-01"),
        (time(8, 0), "08:00"),
    ],
)
def test_dates(value, expected):
    assert display_value(value) == expected


@pytest.mark.parametrize(
    ("value", "fmt", "expected"),
    [
        (123, "00000", "00123"),
        (7, "000", "007"),
        (123456, "00000", "123456"),
        (123, "General", "123"),
        (30000.0, "General", "30000"),
        (0.1 + 0.2, "General", "0.3"),
        (1234.5, "0.00", "1234.50"),
        (2.345, "0.00", "2.35"),  # 四捨五入與 Excel 相同
        (1234567, "#,##0", "1,234,567"),
        (-1234.5, "#,##0.00", "-1,234.50"),
        (0.153, "0%", "15%"),
        (0.1534, "0.0%", "15.3%"),
        (1.5, "0.##", "1.5"),
        (2.0, "0.##", "2"),
        (5, "[Red]0000", "0005"),
        (123, "@", "123"),
        (1.5, "0.0?", "1.5"),  # ? 與 # 相同,只是 Excel 會補空白對齊
        (1.25, "0.0?", "1.25"),
    ],
)
def test_number_formats(value, fmt, expected):
    assert format_number(value, fmt) == expected


@pytest.mark.parametrize(
    ("value", "fmt", "expected"),
    [
        (1.5, "# ?/?", "1 1/2"),
        (0.5, "# ?/?", "1/2"),
        (2, "# ?/?", "2"),
        (0, "# ?/?", "0"),
        (-1.25, "# ?/?", "-1 1/4"),
        (1.6, "# ?/?", "1 3/5"),
        (3.14159, "# ?/?", "3 1/7"),  # 分母最多 1 位數時最接近的分數
        (3.14159, "# ??/??", "3 14/99"),  # 分母最多 2 位數
        (3.14159, "# ???/???", "3 16/113"),
        (0.3333, "# ??/??", "1/3"),
        (0.99, "# ?/?", "1"),  # 進位成整數
        (2.99, "# ?/?", "3"),
        (1.5, "?/?", "3/2"),  # 沒有整數部分時顯示假分數
        (-1.5, "?/?", "-3/2"),
        (0.75, "# ?/4", "3/4"),  # 固定分母
        (1.3, "# ?/4", "1 1/4"),
        (1.5, "# ??/16", "1 8/16"),  # 固定分母不約分
        (1.5, '# ?/? "kg"', "1 1/2 kg"),
        (12345, "0.00E+00", "1.23E+04"),
        (0.000123, "0.00E+00", "1.23E-04"),
        (12345, "0.0E+0", "1.2E+4"),
        (-12345, "0.00E+00", "-1.23E+04"),
        (0, "0.00E+00", "0.00E+00"),
        (99999, "0.00E+00", "1.00E+05"),  # 四捨五入後進位
        (12345, "0.00E-00", "1.23E04"),  # E- 只在指數為負時顯示符號
        (0.00123, "0.00E-00", "1.23E-03"),
        (12345, "00.00E+00", "12.35E+03"),
        (12345, "##0.0E+0", "12.3E+3"),  # 工程記號:指數為 3 的倍數
        (1234567, "##0.0E+0", "1.2E+6"),
        (0.0012345, "##0.0E+0", "1.2E-3"),
        (1e100, "0.00E+00", "1.00E+100"),
        (12345, "0E+00", "1E+04"),
        (12345, "0.##E+00", "1.23E+04"),
        (10000, "0.##E+00", "1E+04"),
    ],
)
def test_fraction_and_scientific_formats(value, fmt, expected):
    assert format_number(value, fmt) == expected


def test_fraction_and_scientific_searchable(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["品名", "尺寸", "數量"])
    ws.append(["螺絲", 1.5, 12345])
    ws["B2"].number_format = "# ?/?"
    ws["C2"].number_format = "0.00E+00"
    wb.save(tmp_path / "格式.xlsx")
    assert excel_reader.read_excel_safely(tmp_path / "格式.xlsx")["Sheet"].rows == [["螺絲", "1 1/2", "1.23E+04"]]
    assert search(tmp_path, "1 1/2")["results"]
    assert search(tmp_path, "1.23E+04")["results"]


@pytest.mark.parametrize(
    ("value", "fmt", "expected"),
    [
        (1000, '"NT$"#,##0', "NT$1,000"),
        (1000, "[$NT$-404]#,##0", "NT$1,000"),
        (1234.5, '[$NT$-404]#,##0.00', "NT$1,234.50"),
        (1000, "$#,##0", "$1,000"),
        (-1000, "$#,##0", "-$1,000"),
        (-1000, "$#,##0;($#,##0)", "($1,000)"),
        (-1000, '"NT$"#,##0;[Red]"NT$"#,##0', "NT$1,000"),  # 負數段沒有負號,Excel 也不顯示
        (1000, '#,##0 "元"', "1,000 元"),
        (1000, "#,##0\\元", "1,000元"),
        (1000, "¥#,##0", "¥1,000"),
        (1000, "€#,##0.00", "€1,000.00"),
        (1000, '_-"NT$"* #,##0_-;-"NT$"* #,##0_-;_-"NT$"* "-"_-;_-@_-', "NT$1,000"),  # 會計格式(* 為填滿字元,省略)
        (0, '_-"NT$"* #,##0_-;-"NT$"* #,##0_-;_-"NT$"* "-"_-;_-@_-', "NT$-"),
        (5, '"NT$"General', "NT$5"),
        (5, '"A;B"0', "A;B5"),  # 引號內的 ; 不是分段
    ],
)
def test_currency_formats(value, fmt, expected):
    assert format_number(value, fmt) == expected


def test_bool_and_none():
    assert display_value(True) == "TRUE"
    assert display_value(None) == ""


def test_normalize():
    assert normalize("2026/01/15") == normalize("2026-01-15")
    assert normalize("1,234,567") == "1234567"
    assert normalize("王,小明") == "王,小明"  # 非數字間的逗號保留
    assert normalize("ABC") == "abc"


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    d = tmp_path / "data"
    d.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["料號", "品名", "簽約日", "到期時間", "金額", "單價"])
    ws.append([123, "螺絲", datetime(2026, 1, 15), datetime(2026, 12, 31, 17, 30), 1234567, 1000])
    ws["A2"].number_format = "00000"
    ws["C2"].number_format = "yyyy/mm/dd"
    ws["D2"].number_format = "yyyy/mm/dd hh:mm"
    ws["E2"].number_format = "#,##0"
    ws["F2"].number_format = '"NT$"#,##0'
    wb.save(d / "料號.xlsx")
    return d


def test_read_restores_display(data_dir):
    sheet = excel_reader.read_excel_safely(data_dir / "料號.xlsx")["Sheet"]
    assert sheet.rows == [["00123", "螺絲", "2026-01-15", "2026-12-31 17:30", "1,234,567", "NT$1,000"]]


@pytest.mark.parametrize(
    "query",
    ["00123", "2026/01/15", "2026-01-15", "2026/12/31 17:30", "1234567", "1,234,567", "NT$1,000", "nt$1000"],
)
def test_search_matches_displayed_text(data_dir, query):
    results = search(data_dir, query)["results"]
    assert results and results[0]["rows"][0][1] == "螺絲"


@pytest.mark.parametrize("query", ["1000", "1,000"])
def test_search_currency_by_number(data_dir, query):
    rows = search(data_dir, query)["results"][0]["rows"]
    assert rows[0][5] == "NT$1,000"


def test_search_does_not_match_across_cells(data_dir):
    # 「螺絲」與「2026」分屬相鄰儲存格,不應被當成連續文字
    assert search(data_dir, "螺絲2026")["results"] == []
