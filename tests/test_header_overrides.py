"""手動指定表頭列(header_overrides):有設定時不做自動偵測;設定錯誤時記錄警告但不當掉。"""

import json
import logging
from pathlib import Path

import pytest
from openpyxl import Workbook

import excel_reader
from conftest import loaded_client

ROOT = Path(__file__).resolve().parents[1]

# 自動偵測會判斷錯誤的表:表頭大多空白,資料列每格都是文字 → 自動偵測把第 3 列(第一筆資料)當表頭
SHEET1 = [
    ["醫材清單(2026 年度)"],
    ["品名", None, None, None],
    ["紗布", "盒", "庫房A", "常備"],
    ["針筒", "支", "庫房B", "常備"],
]
SHEET2 = [
    ["說明:本表每月更新"],
    ["代碼", "名稱"],
    ["M01", "口罩"],
]


def make(path: Path, sheets: dict[str, list[list]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    wb.save(path)


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    d = tmp_path / "data"
    make(d / "醫材.xlsx", {"工作表1": SHEET1, "工作表2": SHEET2})
    return d


def catalog_with(data_dir: Path, overrides: dict) -> excel_reader.Catalog:
    catalog = excel_reader.Catalog(data_dir, header_overrides=overrides)
    catalog.refresh()
    return catalog


def sheet(catalog, name: str, sheet_name: str) -> dict:
    return catalog.rows(name, sheet_name)


def test_auto_detection_is_wrong_without_override(data_dir):
    """前提:沒有設定時,自動偵測把第一筆資料當成表頭。"""
    assert sheet(catalog_with(data_dir, {}), "醫材", "工作表1")["columns"][0] == "紗布"


def test_file_override_applies_to_all_sheets(data_dir):
    catalog = catalog_with(data_dir, {"醫材.xlsx": 2})
    first = sheet(catalog, "醫材", "工作表1")
    assert first["columns"] == ["品名", "欄位2", "欄位3", "欄位4"]
    assert first["rows"] == [["紗布", "盒", "庫房A", "常備"], ["針筒", "支", "庫房B", "常備"]]
    assert sheet(catalog, "醫材", "工作表2")["columns"] == ["代碼", "名稱"]


def test_sheet_override_takes_priority(data_dir):
    catalog = catalog_with(data_dir, {"醫材.xlsx": 3, "醫材.xlsx/工作表1": 2})
    assert sheet(catalog, "醫材", "工作表1")["columns"][0] == "品名"
    second = sheet(catalog, "醫材", "工作表2")  # 套用整個檔案的設定:第 3 列為表頭
    assert second["columns"] == ["M01", "口罩"] and second["rows"] == []


def test_rows_above_header_excluded_and_not_searchable(data_dir):
    catalog = catalog_with(data_dir, {"醫材.xlsx": 2})
    assert catalog.search("2026 年度")["results"] == []
    assert catalog.search("紗布")["results"][0]["rows"] == [["紗布", "盒", "庫房A", "常備"]]


def test_override_counts_hidden_rows(tmp_path):
    """設定值是 Excel 上看到的列號,上方有隱藏列時仍指到同一列;指定的列本身被隱藏也照常當表頭。"""
    d = tmp_path / "data"
    d.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "清單"
    for row in [["內部說明"], ["標題"], ["品名", None, None], ["紗布", "盒", "庫房A"], ["針筒", "支", "庫房B"]]:
        ws.append(row)
    ws.row_dimensions[1].hidden = True
    wb.save(d / "隱藏.xlsx")
    assert sheet(catalog_with(d, {"隱藏.xlsx": 3}), "隱藏", "清單")["columns"] == ["品名", "欄位2", "欄位3"]
    ws.row_dimensions[3].hidden = True
    wb.save(d / "隱藏.xlsx")
    assert sheet(catalog_with(d, {"隱藏.xlsx": 3}), "隱藏", "清單")["columns"] == ["品名", "欄位2", "欄位3"]


def test_override_in_subfolder_and_key_variants(tmp_path):
    d = tmp_path / "data"
    make(d / "醫材" / "耗材.XLSX", {"工作表1": SHEET1})
    for key in ("醫材/耗材.XLSX", "醫材\\耗材.xlsx", "醫材/耗材.xlsx/工作表1"):
        assert sheet(catalog_with(d, {key: 2}), "醫材/耗材", "工作表1")["columns"][0] == "品名", key


def test_missing_file_and_sheet_warn_but_do_not_crash(data_dir, caplog):
    overrides = {"不存在.xlsx": 2, "醫材.xlsx/不存在的工作表": 2, "醫材.xlsx/工作表1": 2}
    with caplog.at_level(logging.WARNING):
        catalog = catalog_with(data_dir, overrides)
        catalog.refresh()  # 第二次檢查不重複記錄
    assert sheet(catalog, "醫材", "工作表1")["columns"][0] == "品名"
    messages = [r.getMessage() for r in caplog.records if "header_overrides" in r.getMessage()]
    assert len([m for m in messages if "不存在.xlsx" in m]) == 1
    assert len([m for m in messages if "不存在的工作表" in m]) == 1


@pytest.mark.parametrize("value", [0, -1, "三", 2.5, True, None])
def test_invalid_value_ignored_with_warning(data_dir, caplog, value):
    with caplog.at_level(logging.WARNING):
        catalog = catalog_with(data_dir, {"醫材.xlsx": value})
    assert sheet(catalog, "醫材", "工作表1")["columns"][0] == "紗布"  # 改用自動偵測
    assert any("header_overrides" in r.getMessage() for r in caplog.records)


def test_row_beyond_sheet_falls_back_with_warning(data_dir, caplog):
    with caplog.at_level(logging.WARNING):
        catalog = catalog_with(data_dir, {"醫材.xlsx/工作表2": 50})
    assert sheet(catalog, "醫材", "工作表2")["columns"] == ["代碼", "名稱"]
    assert any("50" in r.getMessage() and "工作表2" in r.getMessage() for r in caplog.records)


def test_not_a_dict_ignored(data_dir, caplog):
    with caplog.at_level(logging.WARNING):
        catalog = catalog_with(data_dir, ["醫材.xlsx", 2])
    assert catalog.categories()[0]["status"] == "ok"
    assert any("header_overrides" in r.getMessage() for r in caplog.records)


def test_config_option(data_dir):
    client = loaded_client({"data_dir": data_dir, "header_overrides": {"醫材.xlsx/工作表1": 2}})
    res = client.get("/api/files/醫材/rows", params={"sheet": "工作表1"}).json()
    assert res["columns"] == ["品名", "欄位2", "欄位3", "欄位4"]


def test_example_config_has_empty_overrides():
    config = json.loads((ROOT / "config.example.json").read_text(encoding="utf-8"))
    assert config["header_overrides"] == {}


def test_build_table_with_fixed_header():
    from table_layout import build_table

    grid = [["標題", ""], ["a", ""], ["1", "2"]]
    is_text = [[True, False], [True, False], [False, False]]
    assert build_table(grid, is_text, [], header_row=1) == (["a", "欄位2"], [["1", "2"]])
