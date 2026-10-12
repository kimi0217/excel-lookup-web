"""合併儲存格、隱藏列欄、篩選範圍都從工作表 XML 解析;工作表 XML 的位置由 xl/workbook.xml 與其 rels 決定,
不依賴 openpyxl 的內部屬性。"""

import logging
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import excel_reader
from app import create_app
from fixtures.generate_fixtures import merged_cells

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def merged_file(tmp_path: Path) -> Path:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    path = data_dir / "02_合併儲存格.xlsx"
    merged_cells(path)  # 與 tests/fixtures/generated/ 相同的產生方式
    return path


def test_merge_ranges_found_in_fixture(merged_file):
    with zipfile.ZipFile(merged_file) as archive:
        paths = excel_reader.sheet_xml_paths(archive)
        merges = excel_reader._parse_merges(archive.read(paths["部門人員"]))
    assert sorted(merges, key=lambda m: (m.min_row, m.min_col)) == [
        excel_reader.Merge(0, 0, 1, 0),  # A1:A2 表頭「部門」
        excel_reader.Merge(0, 1, 0, 2),  # B1:C1 表頭「人員」
        excel_reader.Merge(2, 0, 4, 0),  # A3:A5 資料「資訊部」
    ]


def test_merged_values_filled_in_fixture(merged_file):
    rows = excel_reader.read_excel_safely(merged_file)["部門人員"].rows
    assert [r[:2] for r in rows if r[1] in ("王小明", "陳大文", "林美玲")] == [
        ["資訊部", "王小明"], ["資訊部", "陳大文"], ["資訊部", "林美玲"],
    ]


def test_no_dependency_on_openpyxl_internal_attribute(merged_file, monkeypatch):
    """模擬 openpyxl 升級後把內部屬性 _worksheet_path 改名:openpyxl 自己照常運作,合併儲存格仍要填值。"""
    from openpyxl.worksheet import _read_only

    original_init = _read_only.ReadOnlyWorksheet.__init__

    def renamed_init(self, parent_workbook, title, worksheet_path, shared_strings):
        original_init(self, parent_workbook, title, worksheet_path, shared_strings)
        self._renamed_path = self.__dict__.pop("_worksheet_path")

    monkeypatch.setattr(_read_only.ReadOnlyWorksheet, "__init__", renamed_init)
    monkeypatch.setattr(  # __init__ 中途就會讀取工作表,當時還沒改名
        _read_only.ReadOnlyWorksheet, "_get_source",
        lambda self: self.parent._archive.open(self.__dict__.get("_renamed_path") or self.__dict__["_worksheet_path"]),
    )
    rows = excel_reader.read_excel_safely(merged_file)["部門人員"].rows
    assert ["資訊部", "陳大文", "102"] in rows


def test_source_does_not_mention_internal_attribute():
    for name in ("excel_reader.py", "app.py"):
        assert "_worksheet_path" not in (ROOT / name).read_text(encoding="utf-8"), name


def swap_sheet_files(path: Path) -> None:
    """把工作表 1、2 的 XML 檔名互換並同步修改 rels:內容意義不變,但「第 1 個工作表 = sheet1.xml」不再成立。
    rId1 改用相對路徑、rId2 用絕對路徑,兩種寫法都要支援。"""
    with zipfile.ZipFile(path) as src:
        items = {name: src.read(name) for name in src.namelist()}
    a, b = "xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"
    items[a], items[b] = items[b], items[a]
    rels = items["xl/_rels/workbook.xml.rels"].decode()
    rels = rels.replace('Target="/xl/worksheets/sheet1.xml" Id="rId1"', 'Target="worksheets/sheet2.xml" Id="rId1"')
    rels = rels.replace('Target="/xl/worksheets/sheet2.xml" Id="rId2"', 'Target="/xl/worksheets/sheet1.xml" Id="rId2"')
    assert "sheet2.xml\" Id=\"rId1\"" in rels and "sheet1.xml\" Id=\"rId2\"" in rels
    items["xl/_rels/workbook.xml.rels"] = rels.encode()
    with zipfile.ZipFile(path, "w") as dst:
        for name, data in items.items():
            dst.writestr(name, data)


def test_sheet_paths_follow_relationships(tmp_path):
    wb = Workbook()
    first = wb.active
    first.title = "研發&製造 <一>"
    first.append(["部門", "姓名"])
    first.append(["研發", "王小明"])
    first.append([None, "陳大文"])
    first.merge_cells("A2:A3")
    second = wb.create_sheet("清單")
    second.append(["編號", "名稱"])
    second.append(["A1", "可見"])
    second.append(["A2", "手動隱藏"])
    second.row_dimensions[3].hidden = True
    path = tmp_path / "rels.xlsx"
    wb.save(path)
    swap_sheet_files(path)

    with zipfile.ZipFile(path) as archive:
        assert excel_reader.sheet_xml_paths(archive) == {
            "研發&製造 <一>": "xl/worksheets/sheet2.xml",
            "清單": "xl/worksheets/sheet1.xml",
        }
    sheets = excel_reader.read_excel_safely(path)
    assert sheets["研發&製造 <一>"].rows == [["研發", "王小明"], ["研發", "陳大文"]]
    assert sheets["清單"].rows == [["A1", "可見"]]


def test_sheet_xml_not_found_does_not_crash(merged_file, monkeypatch, caplog):
    """找不到工作表 XML(檔案結構異常)時不當掉:合併格不填值,但仍讀得到資料,並記錄警告。"""
    monkeypatch.setattr(excel_reader, "sheet_xml_paths", lambda archive: {})
    with caplog.at_level(logging.WARNING):
        rows = excel_reader.read_excel_safely(merged_file)["部門人員"].rows
    assert ["", "陳大文", "102"] in rows
    assert "找不到工作表 XML" in caplog.text


def test_sheet_xml_paths_on_incomplete_archive(tmp_path):
    path = tmp_path / "x.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("xl/workbook.xml", "<workbook><sheets><sheet name='a'")  # XML 不完整
    with zipfile.ZipFile(path) as z:
        assert excel_reader.sheet_xml_paths(z) == {}
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("other.txt", "沒有活頁簿")
    with zipfile.ZipFile(path) as z:
        assert excel_reader.sheet_xml_paths(z) == {}


def test_startup_has_no_merge_warning(tmp_path, caplog):
    with caplog.at_level(logging.WARNING):
        with TestClient(create_app({"data_dir": tmp_path})):
            pass
    assert not any("合併儲存格" in r.getMessage() for r in caplog.records)
    assert not hasattr(excel_reader, "merge_support_ok")
