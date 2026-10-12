"""資料檢視頁的 API:檔案資訊、整份資料分頁、篩選、排序、工作表切換。"""

from datetime import datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import excel_reader
from app import create_app
from cell_format import sort_key
from conftest import loaded_client


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    d = tmp_path / "data"
    d.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "清單"
    ws.append(["編號", "名稱", "金額", "日期"])
    for i in range(1, 251):  # 250 筆,每頁 100 筆共 3 頁
        ws.append([f"A{i:03d}", f"品項{i}", i * 10 if i % 50 else None, datetime(2026, 1 + i % 12, 1 + i % 28)])
        ws.cell(i + 1, 3).number_format = '"NT$"#,##0'
    other = wb.create_sheet("備註")
    other.append(["項目", "內容"])
    other.append(["說明", "第二個工作表"])
    hidden = wb.create_sheet("隱藏")
    hidden.append(["項目"])
    hidden.append(["機密"])
    hidden.sheet_state = "hidden"
    wb.save(d / "醫材.xlsx")
    (d / "損毀.xlsx").write_bytes(b"not a zip")
    return d


@pytest.fixture
def client(data_dir) -> TestClient:
    return loaded_client({"data_dir": data_dir})


def rows(client, **params) -> dict:
    res = client.get("/api/files/醫材/rows", params=params)
    assert res.status_code == 200, res.text
    return res.json()


# ---------- 檔案列表 ----------

def test_categories_include_counts(client):
    cats = {c["name"]: c for c in client.get("/api/categories").json()}
    assert (cats["醫材"]["rows"], cats["醫材"]["sheets"]) == (251, 2)  # 隱藏工作表不算
    assert cats["損毀"]["status"] == "error" and cats["損毀"]["rows"] is None
    assert cats["損毀"]["updated"]  # 讀取失敗也顯示檔案更新時間


def test_file_info(client):
    info = client.get("/api/files/醫材").json()
    assert info["status"] == "ok"
    assert info["sheet_list"] == [
        {"name": "清單", "rows": 250, "columns": 4},
        {"name": "備註", "rows": 1, "columns": 2},
    ]


def test_file_info_not_found(client):
    assert client.get("/api/files/不存在").status_code == 404
    assert client.get("/api/files/..%2F..%2Fetc").status_code == 404


def test_file_info_error(client):
    info = client.get("/api/files/損毀").json()
    assert info["status"] == "error" and info["sheet_list"] == []


# ---------- 整份資料分頁 ----------

def test_default_first_page_without_query(client):
    body = rows(client)
    assert body["sheet"] == "清單" and body["columns"] == ["編號", "名稱", "金額", "日期"]
    assert (body["total"], body["matched"], body["page"], body["pages"], body["offset"]) == (250, 250, 1, 3, 0)
    assert len(body["rows"]) == 100
    assert body["rows"][0] == ["A001", "品項1", "NT$10", "2026-02-02"]


def test_last_page_partial(client):
    body = rows(client, page=3)
    assert len(body["rows"]) == 50 and body["offset"] == 200 and body["rows"][-1][0] == "A250"


def test_page_beyond_range_clamped(client):
    body = rows(client, page=99)
    assert body["page"] == 3 and body["rows"][0][0] == "A201"


def test_page_size_limit(client):
    assert client.get("/api/files/醫材/rows", params={"page_size": 501}).status_code == 422
    assert len(rows(client, page_size=500)["rows"]) == 250


def test_other_sheet(client):
    body = rows(client, sheet="備註")
    assert body["rows"] == [["說明", "第二個工作表"]] and body["total"] == 1


def test_hidden_sheet_not_accessible(client):
    assert client.get("/api/files/醫材/rows", params={"sheet": "隱藏"}).status_code == 404


def test_unknown_sheet(client):
    res = client.get("/api/files/醫材/rows", params={"sheet": "不存在"})
    assert res.status_code == 404 and res.json()["detail"] == "找不到此工作表"


def test_error_file_rows(client):
    body = client.get("/api/files/損毀/rows").json()
    assert body["status"] == "error" and body["rows"] == [] and body["error"] == excel_reader.MSG_CORRUPTED


def test_loading_file_rows(data_dir):
    client = TestClient(create_app({"data_dir": data_dir}))  # 不預先載入
    body = client.get("/api/files/醫材/rows").json()
    assert body["status"] == "loading" and body["rows"] == []


# ---------- 篩選 ----------

def test_filter(client):
    body = rows(client, q="品項12")
    assert body["total"] == 250
    assert [r[1] for r in body["rows"]] == ["品項12", "品項120", "品項121", "品項122", "品項123",
                                            "品項124", "品項125", "品項126", "品項127", "品項128", "品項129"]
    assert body["matched"] == 11 and body["pages"] == 1


def test_filter_uses_search_rules(client):
    # / 與 - 視為相同、忽略千分位
    assert rows(client, q="2026/02/02")["matched"] == rows(client, q="2026-02-02")["matched"] > 0
    assert rows(client, q="1010")["rows"][0][2] == "NT$1,010"


def test_filter_no_match(client):
    body = rows(client, q="不存在的字")
    assert body["matched"] == 0 and body["rows"] == [] and body["pages"] == 1


def test_filter_pages(client):
    body = rows(client, q="品項", page=2)
    assert body["matched"] == 250 and body["offset"] == 100


# ---------- 排序 ----------

def test_sort_numbers_by_value(client):
    body = rows(client, sort=2, page_size=500)
    amounts = [r[2] for r in body["rows"]]
    assert amounts[:3] == ["NT$10", "NT$20", "NT$30"]  # 依數值,不是依字串("NT$100" 不會排在 "NT$20" 前)
    assert amounts[-5:] == [""] * 5  # 空白排最後


def test_sort_descending_blanks_still_last(client):
    amounts = [r[2] for r in rows(client, sort=2, order="desc", page_size=500)["rows"]]
    assert amounts[0] == "NT$2,490" and amounts[-5:] == [""] * 5


def test_sort_text(client):
    names = [r[1] for r in rows(client, sort=1, page_size=500)["rows"]]
    assert names == sorted(names, key=str.casefold)


def test_sort_with_filter_and_page(client):
    body = rows(client, q="品項1", sort=0, order="desc", page=1, page_size=5)
    assert [r[0] for r in body["rows"]] == ["A199", "A198", "A197", "A196", "A195"]
    assert body["matched"] == 1 + 10 + 100  # 品項1、品項10~19、品項100~199


def test_sort_invalid_column(client):
    res = client.get("/api/files/醫材/rows", params={"sort": 9})
    assert res.status_code == 400 and res.json()["detail"] == "排序欄位不存在"
    assert client.get("/api/files/醫材/rows", params={"order": "up"}).status_code == 422


def test_sort_order_cached(data_dir, monkeypatch):
    catalog = excel_reader.Catalog(data_dir)
    catalog.refresh()
    sheet = catalog._entries["醫材"].sheets["清單"]
    calls = []
    real = excel_reader.sort_key
    monkeypatch.setattr(excel_reader, "sort_key", lambda t: calls.append(t) or real(t))
    catalog.rows("醫材", sort=1)
    first = len(calls)
    catalog.rows("醫材", sort=1, descending=True, page=2)
    assert first > 0 and len(calls) == first  # 第二次(含反向)直接用快取
    assert 1 in sheet._sort_cache


@pytest.mark.parametrize(
    ("values", "expected"),
    [
        (["10", "9", "100"], ["9", "10", "100"]),
        (["NT$1,000", "NT$200", "($50)"], ["($50)", "NT$200", "NT$1,000"]),
        (["15%", "5%"], ["5%", "15%"]),
        (["b", "A", "10"], ["10", "A", "b"]),  # 數字排在文字前
        (["2026-10-01", "2026-02-15"], ["2026-02-15", "2026-10-01"]),
    ],
)
def test_sort_key(values, expected):
    assert sorted(values, key=sort_key) == expected


# ---------- 頁面 ----------

def test_index_has_two_page_layout(client):
    html = client.get("/").text
    for text in ("← 返回檔案列表", "搜尋檔名", "搜尋全部檔案", "搜尋資料", "資料為唯讀,如需修改請聯絡資料負責人", "#/file/"):
        assert text in html
