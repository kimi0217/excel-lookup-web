"""表頭偵測、合併儲存格、欄位名稱整理。"""

import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

import excel_reader
from table_layout import Merge, build_table, detect_header, tidy_columns
from conftest import search


def grid_of(rows: list[list]) -> tuple[list[list[str]], list[list[bool]]]:
    """測試用:把 Python 值轉成顯示文字網格與文字旗標(數字不算文字)。"""
    width = max(len(r) for r in rows)
    grid = [[("" if v is None else str(v)) for v in r] + [""] * (width - len(r)) for r in rows]
    is_text = [[isinstance(v, str) and bool(v) for v in r] + [False] * (width - len(r)) for r in rows]
    return grid, is_text


# ---------- 表頭偵測(純邏輯) ----------

def test_header_below_title_rows():
    grid, is_text = grid_of([
        ["2026 年度員工名冊"],
        ["製表日期:2026-10-01", None, "製表人:人資部"],
        [],
        ["工號", "姓名", "部門"],
        ["E001", "王小明", "資訊部"],
        ["E002", "李小華", "業務部"],
    ])
    assert detect_header(grid, is_text, []) == 3
    assert build_table(grid, is_text, []) == (
        ["工號", "姓名", "部門"],
        [["E001", "王小明", "資訊部"], ["E002", "李小華", "業務部"]],
    )


def test_title_merged_across_full_width_is_not_header():
    # 標題列跨整列合併,合併後每格都有值,但只有 1 個不同的值
    grid, is_text = grid_of([
        ["員工名冊", None, None],
        ["工號", "姓名", "部門"],
        ["E001", "王小明", "資訊部"],
    ])
    assert detect_header(grid, is_text, [Merge(0, 0, 0, 2)]) == 1


def test_numeric_rows_are_not_header():
    grid, is_text = grid_of([
        ["銷售統計"],
        [2024, 2025, 2026],
        ["一月", "二月", "三月"],
        [100, 200, 300],
    ])
    # 第 2 列都是數字,不是表頭;第 3 列是文字且下面有資料
    assert detect_header(grid, is_text, []) == 2


def test_header_with_some_blank_names():
    grid, is_text = grid_of([
        ["姓名", "電話", None, "備註"],
        ["王小明", "02-1234", "分機 5", "主要"],
    ])
    assert detect_header(grid, is_text, []) == 0


def test_leading_blank_rows_and_columns_removed():
    grid, is_text = grid_of([
        [],
        [],
        [None, "編號", None, "名稱"],
        [None, "A001", None, "筆電"],
        [],
        [None, "A002", None, "螢幕"],
    ])
    assert build_table(grid, is_text, []) == (["編號", "名稱"], [["A001", "筆電"], ["A002", "螢幕"]])


def test_column_with_data_but_blank_header_kept():
    grid, is_text = grid_of([["姓名", None], ["王小明", "備註內容"]])
    assert build_table(grid, is_text, []) == (["姓名", "欄位2"], [["王小明", "備註內容"]])


@pytest.mark.parametrize(
    "rows",
    [
        [[1, 2, 3], [4, 5, 6]],  # 全是數字,沒有像表頭的列
        [["只有一列", "沒有資料"]],  # 下面沒有資料
        [["單欄標題"], ["值1"], ["值2"]],  # 只有一欄
    ],
)
def test_fallback_to_first_row(rows):
    grid, is_text = grid_of(rows)
    assert detect_header(grid, is_text, []) == 0


def test_empty_sheet():
    assert detect_header([], [], []) is None
    assert build_table([[""], [""]], [[False], [False]], []) == ([], [])


def test_header_only_searches_first_10_rows():
    rows = [["說明文字"]] * 12 + [["工號", "姓名"], ["E001", "王小明"]]
    grid, is_text = grid_of(rows)
    assert detect_header(grid, is_text, []) == 0


# ---------- 合併儲存格 ----------

def test_data_merges_filled():
    grid, is_text = grid_of([
        ["部門", "姓名"],
        ["資訊部", "王小明"],
        [None, "陳大文"],
        [None, "林美玲"],
        ["業務部", "李小華"],
    ])
    columns, rows = build_table(grid, is_text, [Merge(1, 0, 3, 0)])
    assert rows == [["資訊部", "王小明"], ["資訊部", "陳大文"], ["資訊部", "林美玲"], ["業務部", "李小華"]]


def test_horizontal_data_merge_filled():
    grid, is_text = grid_of([["項目", "上午", "下午"], ["值班", "王小明", None]])
    _, rows = build_table(grid, is_text, [Merge(1, 1, 1, 2)])
    assert rows == [["值班", "王小明", "王小明"]]


def test_header_merges_not_filled():
    # 表頭的跨欄合併不處理:合併範圍內其他格的欄名視為空白
    grid, is_text = grid_of([["部門", "人員", None], ["資訊部", "王小明", "101"]])
    columns, _ = build_table(grid, is_text, [Merge(0, 1, 0, 2)])
    assert columns == ["部門", "人員", "欄位3"]


# ---------- 欄位名稱 ----------

def test_tidy_columns():
    assert tidy_columns(["姓名", "電話", "電話", "", "備註", "備註", "電話"]) == [
        "姓名", "電話", "電話(2)", "欄位4", "備註", "備註(2)", "電話(3)",
    ]


def test_tidy_columns_avoids_existing_names():
    assert tidy_columns(["電話", "電話(2)", "電話"]) == ["電話", "電話(2)", "電話(3)"]
    assert tidy_columns(["欄位2", ""]) == ["欄位2", "欄位2(2)"]


# ---------- 實際讀檔 ----------

@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    d = tmp_path / "data"
    d.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "部門人員"
    ws.append(["部門人員名冊"])
    ws.merge_cells("A1:D1")  # 跨整列合併的標題
    ws.append([])
    ws.append([None, "部門", "姓名", "電話", "電話", None])
    ws.append([None, "資訊部", "王小明", "101", "0912", "備註"])
    ws.append([None, None, "陳大文", "102", "0913"])
    ws.append([None, None, "林美玲", "103", "0914"])
    ws.merge_cells("B4:B6")
    ws.append([None, "業務部", "李小華", "201", "0915"])
    wb.save(d / "名冊.xlsx")
    return d


def test_read_file_end_to_end(data_dir):
    sheet = excel_reader.read_excel_safely(data_dir / "名冊.xlsx")["部門人員"]
    assert sheet.columns == ["部門", "姓名", "電話", "電話(2)", "欄位5"]
    assert sheet.rows == [
        ["資訊部", "王小明", "101", "0912", "備註"],
        ["資訊部", "陳大文", "102", "0913", ""],
        ["資訊部", "林美玲", "103", "0914", ""],
        ["業務部", "李小華", "201", "0915", ""],
    ]


def test_merged_department_searchable(data_dir):
    rows = search(data_dir, "資訊部")["results"][0]["rows"]
    assert [r[1] for r in rows] == ["王小明", "陳大文", "林美玲"]


def test_title_not_shown_as_data(data_dir):
    assert search(data_dir, "部門人員名冊")["results"] == []


def test_merge_parsing_with_namespace_prefix():
    class FakeSheet:
        _worksheet_path = "xl/worksheets/sheet1.xml"

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/worksheets/sheet1.xml", '<x:mergeCells count="2"><x:mergeCell ref="B4:B6"/><x:mergeCell ref="AA1"/></x:mergeCells>')
    with zipfile.ZipFile(buf) as z:
        assert excel_reader._merge_ranges(z, FakeSheet()) == [Merge(3, 1, 5, 1), Merge(0, 26, 0, 26)]
