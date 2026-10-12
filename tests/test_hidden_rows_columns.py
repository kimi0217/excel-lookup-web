"""隱藏的列與欄:預設不讀手動隱藏的列與欄,但因篩選而隱藏的列照常讀取。"""

import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

import excel_reader
from conftest import loaded_client
from table_layout import Merge

ROOT = Path(__file__).resolve().parents[1]


def labels(sheet) -> list[str]:
    return [row[1] for row in sheet.rows]


@pytest.fixture
def path(tmp_path: Path) -> Path:
    """第 3 列在自動篩選範圍內被隱藏(篩選造成),第 7 列在範圍外被隱藏(手動),C 欄手動隱藏。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "清單"
    ws.append(["編號", "名稱", "機密欄", "備註"])
    for i, label in enumerate(["一般1", "篩選隱藏", "一般2", "一般3", "一般4", "手動隱藏"], start=1):
        ws.append([f"H{i}", label, f"機密{i}", "備註"])
    ws.column_dimensions["C"].hidden = True
    ws.auto_filter.ref = "A1:D6"
    ws.row_dimensions[3].hidden = True
    ws.row_dimensions[7].hidden = True
    p = tmp_path / "隱藏.xlsx"
    wb.save(p)
    return p


def test_default_skips_manually_hidden(path):
    sheet = excel_reader.read_excel_safely(path)["清單"]
    assert sheet.columns == ["編號", "名稱", "備註"]
    assert labels(sheet) == ["一般1", "篩選隱藏", "一般2", "一般3", "一般4"]


def test_include_hidden_rows_columns(path):
    sheet = excel_reader.read_excel_safely(path, include_hidden_rows_columns=True)["清單"]
    assert sheet.columns == ["編號", "名稱", "機密欄", "備註"]
    assert labels(sheet) == ["一般1", "篩選隱藏", "一般2", "一般3", "一般4", "手動隱藏"]


def test_hidden_rows_not_searchable(path):
    catalog = excel_reader.Catalog(path.parent)
    catalog.refresh()
    assert catalog.search("手動隱藏")["results"] == []
    assert catalog.search("機密1")["results"] == []
    assert catalog.search("篩選隱藏")["results"]


def test_table_filter_rows_kept(tmp_path):
    """「格式化為表格」的篩選範圍記錄在表格 XML,不在工作表的 autoFilter。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["編號", "名稱"])
    for i, label in enumerate(["一般", "表格篩選隱藏", "一般", "表格外手動隱藏"], start=1):
        ws.append([f"T{i}", label])
    ws.add_table(Table(displayName="清單表", ref="A1:B4"))
    ws.row_dimensions[3].hidden = True
    ws.row_dimensions[5].hidden = True
    p = tmp_path / "表格.xlsx"
    wb.save(p)
    sheet = excel_reader.read_excel_safely(p)["Sheet"]
    assert labels(sheet) == ["一般", "表格篩選隱藏", "一般"]


def test_hidden_column_range_and_title_row(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["內部備註"])  # 手動隱藏的說明列
    ws.append(["編號", "內部1", "內部2", "名稱"])
    ws.append(["A1", "x", "y", "筆電"])
    ws.row_dimensions[1].hidden = True
    ws.column_dimensions.group("B", "C", hidden=True)  # <col min="2" max="3" hidden="1">
    p = tmp_path / "範圍.xlsx"
    wb.save(p)
    sheet = excel_reader.read_excel_safely(p)["Sheet"]
    assert sheet.columns == ["編號", "名稱"] and sheet.rows == [["A1", "筆電"]]


def test_merge_across_hidden_row_still_filled(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["部門", "姓名"])
    ws.append(["資訊部", "王小明"])
    ws.append([None, "陳大文"])
    ws.append([None, "林美玲"])
    ws.merge_cells("A2:A4")
    ws.row_dimensions[2].hidden = True  # 合併範圍的第一列(有值的那格)被隱藏
    p = tmp_path / "合併.xlsx"
    wb.save(p)
    assert excel_reader.read_excel_safely(p)["Sheet"].rows == [["資訊部", "陳大文"], ["資訊部", "林美玲"]]


def test_drop_hidden_pure():
    from table_layout import drop_hidden

    grid = [["a", "b", "c"], ["1", "2", "3"], ["4", "5", "6"]]
    is_text = [[True] * 3, [False] * 3, [False] * 3]
    g, t, m = drop_hidden(grid, is_text, [Merge(1, 0, 2, 0)], hidden_rows={1}, hidden_cols={1})
    assert g == [["a", "c"], ["1", "6"]] and t == [[True, True], [False, False]]
    assert m == [Merge(1, 0, 1, 0)]  # 合併範圍對應到剩下的列,值移到新的左上角


def test_config_option(path):
    assert loaded_client({"data_dir": path.parent}).get("/api/files/隱藏/rows").json()["columns"] == ["編號", "名稱", "備註"]
    client = loaded_client({"data_dir": path.parent, "include_hidden_rows_columns": True})
    assert client.get("/api/files/隱藏/rows").json()["columns"] == ["編號", "名稱", "機密欄", "備註"]


def test_example_config_default_false():
    config = json.loads((ROOT / "config.example.json").read_text(encoding="utf-8"))
    assert config["include_hidden_rows_columns"] is False
